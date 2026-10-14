"""Build the five source-backed figures for the automotive research portfolio page.

Run with the local Python environment that has openpyxl installed. The source
workbook and news archive stay outside the web root; only aggregate SVGs are
written to my-site/public/assets/autoResearch.
"""

from __future__ import annotations

from collections import Counter
from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape
import json
import os
import re

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PROJECT = Path(os.environ.get("AUTO_RESEARCH_PROJECT", r"C:\YUFiles\DFiles\ZUEL-Era\04_Projects\2026-04-ML舆情分析"))
WORKBOOK = PROJECT / "吉利.xlsx"
NEWS_ARCHIVE = PROJECT / "吉利大模型舆情分析" / "new" / "geely-fetch" / "results.json"
OUTPUT = ROOT / "my-site" / "public" / "assets" / "autoResearch"
MONTH_RE = re.compile(r"(\d{2})年(\d{1,2})月")
MODEL_RE = re.compile(r"^[A-Z][A-Z0-9-]+$")
FULL_DATE_PATTERNS = (
    re.compile(r"(?<!\d)(20\d{2})([-/.])(\d{1,2})\2(\d{1,2})(?!\d)"),
    re.compile(r"(?<!\d)(20\d{2})年(\d{1,2})月(\d{1,2})日?"),
)
NAVY = "#00355f"
BLUE = "#4472c4"
SKY = "#acd3f0"
ORANGE = "#ed7d31"
INK = "#272b30"
MUTED = "#62666b"
GRID = "#d8dde3"


def month_from_cell(value: object) -> str | None:
    match = MONTH_RE.fullmatch(str(value).strip()) if value is not None else None
    if match is None:
        return None
    return f"20{match.group(1)}-{int(match.group(2)):02d}"


def has_full_date(value: object) -> bool:
    text = str(value or "")
    for pattern in FULL_DATE_PATTERNS:
        match = pattern.search(text)
        if match is None:
            continue
        groups = match.groups()
        year, month, day = (groups[0], groups[2], groups[3]) if len(groups) == 4 else groups
        try:
            date(int(year), int(month), int(day))
            return True
        except ValueError:
            pass
    return False


def read_series(sheet, header_row: int, first_row: int, code_col: int) -> dict:
    headers = [month_from_cell(cell.value) for cell in sheet[header_row]]
    values = {}
    repeated = set()
    for row in sheet.iter_rows(min_row=first_row, values_only=True):
        code = str(row[code_col]).strip() if row[code_col] is not None else ""
        if not MODEL_RE.fullmatch(code):
            continue
        for col, month in enumerate(headers):
            if month is None or col >= len(row):
                continue
            raw = row[col]
            if isinstance(raw, bool) or not isinstance(raw, (int, float)) or raw < 0:
                continue
            key = (code, month)
            if key in values:
                repeated.add(key)
            else:
                values[key] = float(raw)
    for key in repeated:
        values.pop(key, None)
    return values


def research_values() -> dict:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    try:
        planned = read_series(workbook["计划"], 2, 3, 1)
        produced = read_series(workbook["产量"], 2, 3, 1)
    finally:
        workbook.close()
    months = [f"{year}-{month:02d}" for year, lo, hi in [(2024, 3, 12), (2025, 1, 9)] for month in range(lo, hi + 1)]
    shared = {code for code, _ in produced} & {code for code, _ in planned}
    cohort = sorted(code for code in shared if all((code, month) in produced and (code, month) in planned for month in months))
    monthly = []
    for month in months:
        actual = sum(produced[code, month] for code in cohort)
        plan = sum(planned[code, month] for code in cohort)
        monthly.append({"month": month, "actual": actual, "plan": plan, "gap_pct": (plan / actual - 1) * 100 if actual else None})

    archive = json.loads(NEWS_ARCHIVE.read_text(encoding="utf-8"))
    articles = [article for article in archive.values() if isinstance(article, dict) and article.get("url")]
    topics = Counter()
    for key in archive:
        query = key.rsplit("_", 1)[0]
        if any(word in query for word in ("召回", "投诉", "舆情")):
            group = "风险与反馈"
        elif any(word in query for word in ("政策", "战略")):
            group = "政策与战略"
        elif any(word in query for word in ("新车", "智能驾驶")):
            group = "产品与技术"
        elif any(word in query for word in ("销量", "出口", "财报", "竞争")):
            group = "市场与经营"
        else:
            group = "品牌动态"
        topics[group] += 1
    fields = {
        "原文链接": sum(bool(a.get("url")) for a in articles),
        "标题": sum(bool(str(a.get("title") or "").strip()) for a in articles),
        "来源": sum(bool(str(a.get("source") or "").strip()) for a in articles),
        "正文≥200字": sum(len(str(a.get("full_text") or "").strip()) >= 200 for a in articles),
        "完整发布日期": sum(has_full_date(a.get("published_date")) for a in articles),
    }
    return {"cohort_size": len(cohort), "monthly": monthly, "news_count": len(articles), "topics": topics, "fields": fields}


def txt(x, y, value, size=16, color=INK, weight=400, anchor="start", extra="") -> str:
    return f'<text x="{x}" y="{y}" fill="{color}" font-size="{size}" font-weight="{weight}" text-anchor="{anchor}" {extra}>{escape(str(value))}</text>'


def line(x1, y1, x2, y2, color=GRID, width=1, dash="") -> str:
    dash_attr = f' stroke-dasharray="{dash}"' if dash else ""
    return f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="{color}" stroke-width="{width}"{dash_attr}/>'


def rect(x, y, width, height, fill) -> str:
    return f'<rect x="{x}" y="{y}" width="{width}" height="{height}" fill="{fill}"/>'


def svg_document(number: int, title: str, subtitle: str, chart: str, note: str, source: str) -> str:
    font = '"SF Pro Text", "SF Pro", -apple-system, BlinkMacSystemFont, Arial, "PingFang SC", "Microsoft YaHei", sans-serif'
    header = [
        rect(0, 0, 920, 590, "#ffffff"),
        txt(54, 39, f"FIG. {number:02d} / AUTO RESEARCH", 12, MUTED, 600, extra='letter-spacing="2"'),
        line(54, 51, 866, 51, "#c7ccd1", 1.2),
        rect(54, 69, 58, 4, ORANGE),
        txt(54, 114, title, 28, INK, 700),
        txt(54, 145, subtitle, 16, MUTED),
        line(54, 497, 866, 497, "#c7ccd1", 1.1),
        txt(54, 526, note, 13, MUTED),
        txt(54, 553, source, 13, MUTED),
    ]
    return (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 920 590" role="img" aria-labelledby="figure-title figure-desc">'
            f'<title id="figure-title">{escape(title)}</title><desc id="figure-desc">{escape(subtitle + "。" + note + "。" + source)}</desc>'
            f'<style>text{{font-family:{font};font-variant-numeric:tabular-nums}}</style>'
            + "".join(header[:6]) + chart + "".join(header[6:]) + "</svg>\n")


def figure_1(values: dict) -> str:
    monthly = values["monthly"]
    max_v = 160000
    x0, x1, y0, y1 = 109, 842, 457, 183
    parts = []
    for amount in [0, 40000, 80000, 120000, 160000]:
        y = y0 - amount / max_v * (y0 - y1)
        parts.extend([line(x0, y, x1, y, GRID, 1, "4 5"), txt(96, y + 5, f"{amount/10000:.0f}", 13, MUTED, anchor="end")])
    parts.append(txt(54, 184, "万辆", 13, MUTED))
    for index in [0, 3, 6, 9, 12, 15, 18]:
        x = x0 + index * (x1 - x0) / 18
        parts.append(txt(x, 478, monthly[index]["month"], 12, MUTED, anchor="middle"))
    for key, color in [("plan", SKY), ("actual", NAVY)]:
        points = [(x0 + i * (x1 - x0) / 18, y0 - row[key] / max_v * (y0 - y1)) for i, row in enumerate(monthly)]
        path = " ".join(f"{x:.1f},{y:.1f}" for x, y in points)
        parts.append(f'<polyline points="{path}" fill="none" stroke="{color}" stroke-width="3.6" stroke-linejoin="round" stroke-linecap="round"/>')
        parts.extend(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="2.8" fill="{color}"/>' for x, y in points)
    parts.extend([line(661, 174, 689, 174, NAVY, 3), txt(699, 179, "实际产量", 14, MUTED), line(771, 174, 799, 174, SKY, 3), txt(809, 179, "计划", 14, MUTED)])
    return svg_document(1, "固定车型组的月度产量与计划", "2024.03—2025.09 · 21 个全期成对车型 · 单位：万辆", "".join(parts),
                        "注：仅纳入 19 个月均有成对记录的车型；计划不等于模型预测。", "资料来源：项目《吉利.xlsx》；按车型逐月汇总。")


def figure_2(values: dict) -> str:
    rows = [row for row in values["monthly"] if row["month"].startswith("2025-")]
    max_pct = 40
    parts = []
    x0, x1, y0, y1 = 111, 841, 439, 201
    for pct in [0, 10, 20, 30, 40]:
        y = y0 - pct / max_pct * (y0 - y1)
        parts.extend([line(x0, y, x1, y, GRID, 1, "4 5"), txt(98, y + 5, f"{pct}%", 13, MUTED, anchor="end")])
    peak_gap = max(row["gap_pct"] for row in rows)
    for i, row in enumerate(rows):
        x = x0 + i * (x1 - x0) / len(rows) + 21
        v = row["gap_pct"]
        h = v / max_pct * (y0 - y1)
        parts.extend([rect(x, y0-h, 47, h, ORANGE if v == peak_gap else NAVY),
                      txt(x+23.5, y0-h-9, f"{v:.1f}%", 13, INK, 600, "middle"),
                      txt(x+23.5, 465, row["month"][-2:], 13, MUTED, anchor="middle")])
    parts.append(txt(54, 184, "(计划−实际) / 实际", 13, MUTED))
    return svg_document(2, "计划产量相对实际产量的偏离", "2025 年 1—9 月 · 同一组 21 个车型 · 按月汇总", "".join(parts),
                        "注：正值表示计划高于实际；分母为当月实际产量之和。", "资料来源：项目《吉利.xlsx》；作者计算。")


def figure_3(values: dict) -> str:
    groups = [("市场与经营", 38), ("品牌动态", 26), ("风险与反馈", 24), ("产品与技术", 19), ("政策与战略", 17)]
    assert dict(groups) == dict(values["topics"])
    parts = [txt(54, 188, "条", 13, MUTED)]
    x0, x1 = 228, 794
    for tick in [0, 10, 20, 30, 40]:
        x = x0 + tick / 40 * (x1 - x0)
        parts.extend([line(x, 196, x, 448, GRID, 1, "4 5"), txt(x, 472, tick, 13, MUTED, anchor="middle")])
    for i, (name, count) in enumerate(groups):
        y = 216 + i * 47
        parts.extend([txt(210, y+20, name, 16, INK, 500, "end"), rect(x0, y, count/40*(x1-x0), 27, NAVY if i != 2 else ORANGE), txt(x0+count/40*(x1-x0)+13, y+20, count, 15, INK, 700)])
    return svg_document(3, "新闻检索任务的主题分布", "项目归档的 124 条检索记录 · 按检索词归类", "".join(parts),
                        "注：反映采样设计，并非社会舆情或文章主题占比。", "资料来源：项目 results.json；每个检索键对应一条记录。")


def figure_4(values: dict) -> str:
    fields = list(values["fields"].items())
    total = values["news_count"]
    parts = [txt(54, 188, "条", 13, MUTED)]
    x0, x1 = 228, 794
    for tick in [0, 31, 62, 93, 124]:
        x = x0 + tick / total * (x1-x0)
        parts.extend([line(x, 195, x, 449, GRID, 1, "4 5"), txt(x, 473, tick, 13, MUTED, anchor="middle")])
    for i, (name, count) in enumerate(fields):
        y = 213 + i * 47
        parts.extend([txt(210, y+21, name, 16, INK, 500, "end"), rect(x0, y, (x1-x0), 28, "#eef3f7"), rect(x0, y, count/total*(x1-x0), 28, ORANGE if name == "完整发布日期" else BLUE), txt(805, y+21, f"{count}/{total}", 15, INK, 700)])
    return svg_document(4, "新闻样本的字段可用性", "项目归档的 124 条 URL 唯一样本", "".join(parts),
                        "注：完整发布日期指含年、月、日；正文以不少于 200 字判定。", "资料来源：项目 results.json；原始字段完整性检查。")


def figure_5() -> str:
    data = [("2024", 1316.8, SKY), ("2025", 1652.4, NAVY)]
    parts = [txt(54, 187, "万辆", 13, MUTED)]
    x0, x1, y0, y1 = 137, 816, 450, 205
    for tick in [0, 500, 1000, 1500, 2000]:
        y = y0 - tick/2000*(y0-y1)
        parts.extend([line(x0, y, x1, y, GRID, 1, "4 5"), txt(124, y+5, tick, 13, MUTED, anchor="end")])
    for i, (year, amount, color) in enumerate(data):
        x = 261 + i*293
        h = amount/2000*(y0-y1)
        parts.extend([rect(x, y0-h, 120, h, color), txt(x+60, y0-h-13, f"{amount:,.1f}", 23, NAVY, 700, "middle"), txt(x+60, 477, year, 16, MUTED, anchor="middle")])
    return svg_document(5, "全国新能源汽车产量：年度参照", "2024—2025 年 · 国家统计局 · 单位：万辆", "".join(parts),
                        "注：全国行业口径，仅作外部背景，不与项目车型组求占比。", "资料来源：国家统计局 2024、2025 年国民经济和社会发展统计公报。")


def build_figures() -> None:
    values = research_values()
    OUTPUT.mkdir(parents=True, exist_ok=True)
    figures = [figure_1(values), figure_2(values), figure_3(values), figure_4(values), figure_5()]
    for index, figure in enumerate(figures, 1):
        (OUTPUT / f"figure-{index:02d}.svg").write_text(figure, encoding="utf-8")
    print(f"Built {len(figures)} figures in {OUTPUT}")
    print(json.dumps({"cohort_size": values["cohort_size"], "news_count": values["news_count"], "topics": values["topics"], "fields": values["fields"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    build_figures()
